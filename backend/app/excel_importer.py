import os
import re
import openpyxl
from sqlalchemy.orm import Session
from app.models import Category, Product, OrderItem

DEFAULT_EXCEL_PATH = r"C:\Users\VASANTH A\OneDrive\Documents\Bermuda_Menu PRICE LIST NEW_Excel.xlsx"

CATEGORY_ICONS = {
    'SOUPS': '🍲',
    'VEG STARTERS': '🥗',
    'NON VEG STARTERS': '🍗',
    'PLATTERS': '🍱',
    'TANDOOR & GRILL': '🔥',
    'RICE / NOODLES': '🍜',
    'INDIAN BREADS': '🫓',
    'INDIAN CURRIES & GRAVIES': '🥘',
    'PASTA': '🍝',
    'PIZZA': '🍕',
    'BIRYANI & PULAO': '🍚',
    'DESSERTS': '🍰',
    'CLASSIC COCKTAILS': '🍸',
    'BERMUDA SIGNATURE COCKTAILS': '🍹',
    'BERMUDA LONGDRIVE SIGNATURE COCKTAILS': '🍸',
    'LONG DRIVE PITCHER COCKTAIL': '🍹',
    'SHOOTERS': '🥃',
    'SINGLE MALT': '🥃',
    'SCOTCH PREMIUM / INTERNATIONAL WHISKEY': '🥃',
    'REGULAR WHISKEY': '🥃',
    'GIN': '🍸',
    'VODKA': '🍸',
    'RUM': '🥃',
    'BRANDY': '🥃',
    'LIQUEURS': '🍾',
    'WINE': '🍷',
    'TEQUILA': '🌵',
    'DOMESTIC BEER': '🍺',
    'IMPORTED': '🍻',
    'ZERO ALCOHOL': '🥤',
    'SOFT DRINKS & BEVERAGES': '🥤'
}

def import_excel_menu(db: Session, excel_path: str = DEFAULT_EXCEL_PATH):
    if not os.path.exists(excel_path):
        print(f"[excel_importer] File not found: {excel_path}")
        return False

    try:
        wb = openpyxl.load_workbook(excel_path, data_only=True)
    except Exception as e:
        print(f"[excel_importer] Failed to load excel workbook: {e}")
        return False
    
    categories_map = {} # cat_name -> {target_dept, icon}
    products_list = []  # list of {name, cat_name, price, target_dept, desc}

    # 1. KITCHEN Sheet
    if 'KITCHEN' in wb.sheetnames:
        ws = wb['KITCHEN']
        for row in list(ws.iter_rows(values_only=True))[1:]:
            if not any(row):
                continue
            sec, item_name, veg_type, desc, price = row[0], row[1], row[2], row[3], row[4]
            if not item_name or not str(item_name).strip():
                continue
            
            cat_name = str(sec or 'KITCHEN STARTERS & MAINS').strip().upper()
            cat_name = re.sub(r'[\:\s]+$', '', cat_name)
            
            if cat_name not in categories_map:
                categories_map[cat_name] = {'target_dept': 'KITCHEN', 'icon': CATEGORY_ICONS.get(cat_name, '🍳')}
                
            price_str = str(price).strip() if price is not None else '499'
            item_base = str(item_name).strip()
            desc_str = str(desc).strip() if desc else None

            # Split prices like 349/399 or 499/549/599 or 599/999
            if '/' in price_str and not price_str.startswith('='):
                parts = [p.strip() for p in price_str.split('/')]
                if 'half / full' in item_base.lower() or 'half/full' in item_base.lower():
                    base_clean = re.sub(r'(?i)\s*[\-\–]?\s*half\s*\/\s*full', '', item_base).strip()
                    if len(parts) >= 2:
                        products_list.append({'name': f'{base_clean} (Half)', 'cat': cat_name, 'price': float(parts[0]), 'target_dept': 'KITCHEN', 'desc': desc_str})
                        products_list.append({'name': f'{base_clean} (Full)', 'cat': cat_name, 'price': float(parts[1]), 'target_dept': 'KITCHEN', 'desc': desc_str})
                elif len(parts) == 2:
                    products_list.append({'name': f'{item_base} (Veg)', 'cat': cat_name, 'price': float(parts[0]), 'target_dept': 'KITCHEN', 'desc': desc_str})
                    products_list.append({'name': f'{item_base} (Non-Veg)', 'cat': cat_name, 'price': float(parts[1]), 'target_dept': 'KITCHEN', 'desc': desc_str})
                elif len(parts) == 3:
                    variants = ['Veg', 'Chicken', 'Prawn / Paneer']
                    for i, p_val in enumerate(parts):
                        v_label = variants[i] if i < len(variants) else f'Variant {i+1}'
                        try:
                            pv = float(p_val)
                        except ValueError:
                            pv = 499.0
                        products_list.append({'name': f'{item_base} ({v_label})', 'cat': cat_name, 'price': pv, 'target_dept': 'KITCHEN', 'desc': desc_str})
            else:
                try:
                    p_val = float(price_str)
                except ValueError:
                    p_val = 499.0
                products_list.append({'name': item_base, 'cat': cat_name, 'price': p_val, 'target_dept': 'KITCHEN', 'desc': desc_str})

    # 2. BAR Sheet
    if 'BAR' in wb.sheetnames:
        ws = wb['BAR']
        curr_cat = 'CLASSIC COCKTAILS'
        for row in list(ws.iter_rows(values_only=True))[1:]:
            if not any(row):
                continue
            sec, item_name, desc, price = row[0], row[1], row[2], row[3]
            if sec and str(sec).strip():
                curr_cat = str(sec).strip().upper()
                curr_cat = re.sub(r'[\:\s]+$', '', curr_cat)
            if not item_name or not str(item_name).strip():
                continue

            if curr_cat not in categories_map:
                categories_map[curr_cat] = {'target_dept': 'BAR', 'icon': CATEGORY_ICONS.get(curr_cat, '🍸')}

            desc_str = str(desc).strip() if desc else None
            try:
                p_val = float(price) if price is not None else 499.0
            except (ValueError, TypeError):
                p_val = 499.0

            products_list.append({'name': str(item_name).strip(), 'cat': curr_cat, 'price': p_val, 'target_dept': 'BAR', 'desc': desc_str})

    # 3. SOFT DRINKS Sheet
    if 'SOFT DRINKS' in wb.sheetnames:
        ws = wb['SOFT DRINKS']
        soft_cat = 'SOFT DRINKS & BEVERAGES'
        if soft_cat not in categories_map:
            categories_map[soft_cat] = {'target_dept': 'BAR', 'icon': '🥤'}

        for row in list(ws.iter_rows(values_only=True)):
            if not any(row):
                continue
            item_n = None
            price_v = None
            for cell in row:
                if cell is None:
                    continue
                cell_str = str(cell).strip()
                if not cell_str or cell_str == '=':
                    continue
                if re.match(r'^\d+(\.\d+)?$', cell_str):
                    price_v = float(cell_str)
                elif not item_n and len(cell_str) > 1 and not cell_str.startswith('='):
                    item_n = cell_str

            if item_n:
                if not price_v:
                    price_v = 149.0
                products_list.append({'name': item_n, 'cat': soft_cat, 'price': price_v, 'target_dept': 'BAR', 'desc': None})

    # Delete OrderItem references to existing products to avoid foreign key issues
    db.query(OrderItem).delete(synchronize_session=False)
    db.query(Product).delete(synchronize_session=False)
    db.query(Category).delete(synchronize_session=False)
    db.commit()

    # Insert new categories
    created_categories = {}
    for cat_name, cat_info in categories_map.items():
        cat_obj = Category(name=cat_name, target_dept=cat_info['target_dept'], icon=cat_info['icon'])
        db.add(cat_obj)
        db.flush()
        created_categories[cat_name] = cat_obj.id

    # Insert new products
    products_to_add = []
    for prod in products_list:
        cat_id = created_categories.get(prod['cat'])
        if cat_id:
            products_to_add.append(Product(
                name=prod['name'],
                category_id=cat_id,
                price=prod['price'],
                description=prod['desc'],
                target_dept=prod['target_dept'],
                is_available=True
            ))

    db.add_all(products_to_add)
    db.commit()
    print(f"[excel_importer] Successfully imported {len(categories_map)} categories and {len(products_to_add)} products from {excel_path}!")
    return True
